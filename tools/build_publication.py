from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from passkeytransit import __version__


TITLE = "Preservación semántica en migraciones de passkeys con CXF y CXP"
DISPLAY_VERSION = "v" + __version__.rsplit(".", 1)[0]
SUBTITLE = f"PasskeyTransit {DISPLAY_VERSION} Informe reproducible de controles e interoperabilidad"


def _plain(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    return text.replace("–", "-")


def _heading(text: str) -> str:
    return re.sub(r"^\d+(?:\.\d+)*\.\s*", "", _plain(text)).rstrip(":. ")


def parse_markdown(path: Path) -> list[tuple[str, object]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks: list[tuple[str, object]] = []
    paragraph: list[str] = []
    index = 1  # source title is replaced with the publication title

    def flush() -> None:
        if paragraph:
            blocks.append(("paragraph", " ".join(part.strip() for part in paragraph)))
            paragraph.clear()

    while index < len(lines):
        line = lines[index]
        if not line.strip():
            flush()
            index += 1
            continue
        if line.startswith("## "):
            flush()
            blocks.append(("heading1", _heading(line[3:])))
            index += 1
            continue
        if line.startswith("### "):
            flush()
            blocks.append(("heading2", _heading(line[4:])))
            index += 1
            continue
        if line.startswith("| "):
            flush()
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            rows = [[_plain(cell.strip()) for cell in row.strip("|").split("|")] for row in table_lines]
            if len(rows) > 1 and all(set(cell) <= {"-", ":"} for cell in rows[1]):
                rows.pop(1)
            blocks.append(("table", rows))
            continue
        if line.startswith("- "):
            flush()
            index += 1
            parts = [line[2:]]
            while index < len(lines) and lines[index].strip() and not lines[index].startswith(("- ", "##", "|")):
                parts.append(lines[index].strip())
                index += 1
            blocks.append(("bullet", _plain(" ".join(parts))))
            continue
        if re.match(r"^\d+\.\s", line):
            flush()
            index += 1
            parts = [line]
            while index < len(lines) and lines[index].strip() and not re.match(r"^\d+\.\s", lines[index]) and not lines[index].startswith("##"):
                parts.append(lines[index].strip())
                index += 1
            blocks.append(("numbered", _plain(" ".join(parts))))
            continue
        paragraph.append(line)
        index += 1
    flush()
    return blocks


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, value: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _keep_table_together(table) -> None:
    for row_index, row in enumerate(table.rows[:-1]):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def build_docx(blocks: list[tuple[str, object]], output: Path) -> None:
    doc = Document()
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Informe reproducible de controles sintéticos e interoperabilidad para migración de passkeys"
    doc.core_properties.author = "PasskeyTransit research team"
    doc.core_properties.keywords = "passkeys, WebAuthn, CXF, CXP, semantic preservation"
    doc.core_properties.comments = "Generated from the versioned PasskeyTransit manuscript"
    doc.core_properties.created = datetime.now(timezone.utc)
    doc.core_properties.modified = datetime.now(timezone.utc)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.9)

    styles = doc.styles
    for name, size in (("Normal", 10.8), ("Title", 24), ("Heading 1", 16), ("Heading 2", 12.5)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 1"].paragraph_format.line_spacing = Pt(20)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)
    styles["Heading 2"].paragraph_format.line_spacing = Pt(15)
    styles["Heading 2"].paragraph_format.keep_with_next = True
    styles["List Bullet"].font.name = "Arial"
    styles["List Bullet"].font.size = Pt(10.8)
    styles["List Bullet"].paragraph_format.left_indent = Inches(0.25)
    styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.18)
    styles["List Bullet"].paragraph_format.space_after = Pt(3)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    border = title_ppr.find(qn("w:pBdr"))
    if border is not None:
        title_ppr.remove(border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run(TITLE)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    run = subtitle.add_run(SUBTITLE)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(65, 78, 92)

    for kind, content in blocks:
        if kind == "heading1":
            if str(content) == "Referencias":
                doc.add_page_break()
            doc.add_paragraph(str(content), style="Heading 1")
        elif kind == "heading2":
            doc.add_paragraph(str(content), style="Heading 2")
        elif kind == "bullet":
            doc.add_paragraph(str(content), style="List Bullet")
        elif kind == "numbered":
            paragraph = doc.add_paragraph(_plain(str(content)))
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            paragraph.paragraph_format.space_after = Pt(3)
        elif kind == "paragraph":
            doc.add_paragraph(_plain(str(content)))
        elif kind == "table":
            rows = content
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.autofit = True
            _set_repeat_table_header(table.rows[0])
            for row_index, row in enumerate(rows):
                for col_index, value in enumerate(row):
                    cell = table.cell(row_index, col_index)
                    cell.text = value
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    _set_cell_margins(cell)
                    if row_index == 0:
                        _set_cell_shading(cell, "23395D")
                        for run in cell.paragraphs[0].runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                    elif row_index % 2 == 0:
                        _set_cell_shading(cell, "EEF3F8")
                    if col_index > 0:
                        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _keep_table_together(table)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(f"PasskeyTransit {DISPLAY_VERSION}")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(90, 90, 90)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def build_pdf(blocks: list[tuple[str, object]], output: Path) -> None:
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="PaperTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=colors.black, alignment=TA_CENTER, spaceAfter=12))
    styles.add(ParagraphStyle(name="PaperSubtitle", parent=styles["Normal"], fontName="Helvetica", fontSize=10, leading=14, textColor=colors.HexColor("#41505F"), alignment=TA_CENTER, spaceAfter=24))
    styles["BodyText"].fontName = "Helvetica"
    styles["BodyText"].fontSize = 9.5
    styles["BodyText"].leading = 13
    styles["BodyText"].spaceAfter = 7
    styles["Heading1"].fontName = "Helvetica-Bold"
    styles["Heading1"].fontSize = 15
    styles["Heading1"].leading = 18
    styles["Heading1"].textColor = colors.black
    styles["Heading1"].spaceBefore = 12
    styles["Heading1"].spaceAfter = 6
    styles["Heading1"].keepWithNext = 1
    styles["Heading2"].fontName = "Helvetica-Bold"
    styles["Heading2"].fontSize = 11.5
    styles["Heading2"].leading = 14
    styles["Heading2"].textColor = colors.black
    styles["Heading2"].spaceBefore = 8
    styles["Heading2"].spaceAfter = 4
    styles["Heading2"].keepWithNext = 1
    styles.add(ParagraphStyle(name="PaperBullet", parent=styles["BodyText"], leftIndent=14, firstLineIndent=-9, spaceAfter=4))
    styles.add(ParagraphStyle(name="PaperReference", parent=styles["BodyText"], leftIndent=18, firstLineIndent=-18, spaceAfter=4))
    styles.add(ParagraphStyle(name="TableHeader", parent=styles["BodyText"], fontName="Helvetica-Bold", textColor=colors.white))

    story = [Paragraph(TITLE, styles["PaperTitle"]), Paragraph(SUBTITLE, styles["PaperSubtitle"])]
    for kind, content in blocks:
        if kind == "heading1":
            story.append(Paragraph(str(content), styles["Heading1"]))
        elif kind == "heading2":
            story.append(Paragraph(str(content), styles["Heading2"]))
            story.append(Spacer(1, 3))
        elif kind == "paragraph":
            story.append(Paragraph(_plain(str(content)), styles["BodyText"]))
        elif kind == "bullet":
            story.append(Paragraph("- " + _plain(str(content)), styles["PaperBullet"]))
        elif kind == "numbered":
            story.append(Paragraph(_plain(str(content)), styles["PaperReference"]))
        elif kind == "table":
            rows = [
                [
                    Paragraph(cell, styles["TableHeader"] if row_index == 0 else styles["BodyText"])
                    for cell in row
                ]
                for row_index, row in enumerate(content)
            ]
            table = Table(rows, colWidths=[2.35 * inch, 0.85 * inch, 1.25 * inch], repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#23395D")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (1, 1), (-1, -1), "CENTER"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF3F8")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.extend([table, Spacer(1, 8)])

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawCentredString(LETTER[0] / 2, 0.45 * inch, f"PasskeyTransit {DISPLAY_VERSION}  |  {doc.page}")
        canvas.restoreState()

    output.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(str(output), pagesize=LETTER, rightMargin=0.9 * inch, leftMargin=0.9 * inch, topMargin=0.75 * inch, bottomMargin=0.7 * inch, title=TITLE, author="PasskeyTransit research team")
    document.build(story, onFirstPage=footer, onLaterPages=footer)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--pdf", type=Path, required=True)
    args = parser.parse_args()
    blocks = parse_markdown(args.source)
    build_docx(blocks, args.docx)
    build_pdf(blocks, args.pdf)
    manifest_path = args.source.parent / "analysis_manifest.json"
    if manifest_path.is_file():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        outputs = manifest.setdefault("output_hashes", {})
        for output in (args.docx, args.pdf):
            outputs[output.name] = hashlib.sha256(output.read_bytes()).hexdigest()
        manifest_path.write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8", newline="\n")


if __name__ == "__main__":
    main()
